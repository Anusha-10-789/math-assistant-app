import io

from docx import Document
from docx.shared import RGBColor

from models import LessonContent
from visual_text import describe_visual

BLACK = RGBColor(0, 0, 0)


def _add_option_paragraph(document: Document, letter: str, content: str) -> None:
    paragraph = document.add_paragraph()
    letter_run = paragraph.add_run(f"{letter})")
    letter_run.bold = True
    letter_run.font.color.rgb = BLACK
    content_run = paragraph.add_run(f" {content}")
    content_run.font.color.rgb = BLACK


def _add_plain_paragraph(document: Document, text: str) -> None:
    run = document.add_paragraph().add_run(text)
    run.font.color.rgb = BLACK


def build_lesson_docx(lesson: LessonContent) -> io.BytesIO:
    document = Document()

    title = document.add_heading(level=1)
    title_run = title.add_run(f"{lesson.topic} — Grade {lesson.grade}")
    title_run.font.color.rgb = BLACK

    concept_heading = document.add_heading(level=2)
    concept_heading_run = concept_heading.add_run("Concept Explanation")
    concept_heading_run.font.color.rgb = BLACK
    _add_plain_paragraph(document, lesson.concept_explanation)

    lecture_heading = document.add_heading(level=2)
    lecture_heading_run = lecture_heading.add_run("Lecture")
    lecture_heading_run.font.color.rgb = BLACK

    for slide in lesson.lecture_slides:
        slide_heading = document.add_heading(level=3)
        slide_heading_run = slide_heading.add_run(slide.title)
        slide_heading_run.font.color.rgb = BLACK
        _add_plain_paragraph(document, slide.content)
        visual_text = describe_visual(slide.visual)
        if visual_text:
            _add_plain_paragraph(document, visual_text)

    questions_heading = document.add_heading(level=2)
    questions_heading_run = questions_heading.add_run("Practice Questions")
    questions_heading_run.font.color.rgb = BLACK

    for mcq in lesson.mcqs:
        topic_paragraph = document.add_heading(level=2)
        topic_run = topic_paragraph.add_run(f"Q.{mcq.question_number}) {mcq.topic}:")
        topic_run.bold = False
        topic_run.font.color.rgb = BLACK

        _add_plain_paragraph(document, mcq.question)

        _add_option_paragraph(document, "A", mcq.option_a)
        _add_option_paragraph(document, "B", mcq.option_b)
        _add_option_paragraph(document, "C", mcq.option_c)
        _add_option_paragraph(document, "D", mcq.option_d)

        _add_plain_paragraph(document, f"Answer: {mcq.correct_answer}")
        _add_plain_paragraph(document, f"Trick: {mcq.trick}")

        visual_text = describe_visual(mcq.visual)
        if visual_text:
            _add_plain_paragraph(document, visual_text)

        _add_plain_paragraph(document, "Explanation:")
        _add_plain_paragraph(document, mcq.explanation)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
