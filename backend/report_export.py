import io
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

from models import QuizResultData

BLACK = RGBColor(0, 0, 0)


def _percent(result: QuizResultData) -> int:
    if result.total == 0:
        return 0
    return round((result.score / result.total) * 100)


def _encouragement(percent: int) -> str:
    if percent >= 80:
        return "Excellent work! You really know this topic."
    if percent >= 50:
        return "Good effort! A little more practice and you'll master this."
    return "Nice try! Let's keep practicing this topic together."


def build_report_docx(topic: str, grade: int, result: QuizResultData) -> io.BytesIO:
    document = Document()
    percent = _percent(result)

    title = document.add_heading(level=1)
    title_run = title.add_run(f"Test Report — {topic} (Grade {grade})")
    title_run.font.color.rgb = BLACK

    score_run = document.add_paragraph().add_run(f"Score: {result.score} / {result.total} ({percent}%)")
    score_run.bold = True
    score_run.font.color.rgb = BLACK

    encouragement_run = document.add_paragraph().add_run(_encouragement(percent))
    encouragement_run.font.color.rgb = BLACK

    if result.missed:
        review_heading = document.add_heading(level=2)
        review_heading_run = review_heading.add_run("Questions to Review")
        review_heading_run.font.color.rgb = BLACK

        for item in result.missed:
            q_paragraph = document.add_paragraph()
            q_run = q_paragraph.add_run(f"Q.{item.question_number}) {item.topic}: {item.question}")
            q_run.font.color.rgb = BLACK

            your_answer_display = (
                f"{item.your_answer}) {item.your_answer_text}" if item.your_answer else "(skipped)"
            )
            correct_answer_display = f"{item.correct_answer}) {item.correct_answer_text}"

            your_run = document.add_paragraph().add_run(f"Your answer: {your_answer_display}")
            your_run.font.color.rgb = BLACK
            correct_run = document.add_paragraph().add_run(f"Correct answer: {correct_answer_display}")
            correct_run.font.color.rgb = BLACK

            if item.explanation:
                explanation_run = document.add_paragraph().add_run(item.explanation)
                explanation_run.font.color.rgb = BLACK

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


_styles = getSampleStyleSheet()
TITLE_STYLE = ParagraphStyle("ReportTitle", parent=_styles["Heading1"], spaceAfter=12)
SECTION_STYLE = ParagraphStyle("ReportSection", parent=_styles["Heading2"], spaceAfter=8)
QUESTION_STYLE = ParagraphStyle("ReportQuestion", parent=_styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4)
BODY_STYLE = ParagraphStyle("ReportBody", parent=_styles["BodyText"], spaceAfter=6)


def _p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text), style)


def build_report_pdf(topic: str, grade: int, result: QuizResultData) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
    )
    percent = _percent(result)

    story = [
        _p(f"Test Report — {topic} (Grade {grade})", TITLE_STYLE),
        _p(f"Score: {result.score} / {result.total} ({percent}%)", BODY_STYLE),
        _p(_encouragement(percent), BODY_STYLE),
        Spacer(1, 10),
    ]

    if result.missed:
        story.append(_p("Questions to Review", SECTION_STYLE))
        for item in result.missed:
            your_answer_display = (
                f"{item.your_answer}) {item.your_answer_text}" if item.your_answer else "(skipped)"
            )
            correct_answer_display = f"{item.correct_answer}) {item.correct_answer_text}"
            block = [
                _p(f"Q.{item.question_number}) {item.topic}: {item.question}", QUESTION_STYLE),
                _p(f"Your answer: {your_answer_display}", BODY_STYLE),
                _p(f"Correct answer: {correct_answer_display}", BODY_STYLE),
            ]
            if item.explanation:
                block.append(_p(item.explanation, BODY_STYLE))
            story.append(KeepTogether(block))

    doc.build(story)
    buffer.seek(0)
    return buffer
